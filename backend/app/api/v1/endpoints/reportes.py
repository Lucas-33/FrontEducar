"""Exportación de reportes a Excel (.xlsx).

Estos endpoints se invocan desde el navegador (botones "Exportar Excel"), por
lo que el JWT se recibe como parámetro de consulta `?token=...` en lugar del
encabezado Authorization.
"""

from io import BytesIO

from fastapi import APIRouter, Depends, HTTPException, Query, status
from fastapi.responses import StreamingResponse
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from sqlalchemy.orm import Session

from backend.app.core.database import get_db
from backend.app.models.seguridad import Usuario
from backend.app.repositories import (
    alumno_repository,
    curso_repository,
    grupo_deporte_repository,
    inscripcion_comedor_repository,
    inscripcion_transporte_repository,
    materia_repository,
    nivel_repository,
    profesor_repository,
)
from backend.app.services.academico import academico_service
from backend.app.services.comunidad import comunidad_service
from backend.app.services.seguridad import seguridad_service
from backend.app.services.servicios import servicios_escolares_service

router = APIRouter()

ENCABEZADO_FILL = PatternFill(
    start_color="0D6E6E", end_color="0D6E6E", fill_type="solid"
)
ENCABEZADO_FONT = Font(color="FFFFFF", bold=True)


def _autorizar_admin(db: Session, token: str | None) -> Usuario:
    if not token:
        raise HTTPException(
            status_code=status.HTTP_401_UNAUTHORIZED,
            detail="Falta el token de autenticación",
        )
    usuario = seguridad_service.get_usuario_from_token(db, token)
    if seguridad_service.get_rol_nombre(db, usuario) != "ADMIN":
        raise HTTPException(
            status_code=status.HTTP_403_FORBIDDEN,
            detail="Solo los administradores pueden exportar reportes",
        )
    return usuario


def _agregar_hoja(wb: Workbook, titulo: str, encabezados: list[str], filas: list[list]):
    ws = wb.create_sheet(title=titulo[:31])
    ws.append(encabezados)
    for celda in ws[1]:
        celda.fill = ENCABEZADO_FILL
        celda.font = ENCABEZADO_FONT
        celda.alignment = Alignment(horizontal="center")
    for fila in filas:
        ws.append(fila)
    for indice, _ in enumerate(encabezados, start=1):
        letra = get_column_letter(indice)
        ancho = max(
            (
                len(str(ws.cell(row=fila, column=indice).value or ""))
                for fila in range(1, ws.max_row + 1)
            ),
            default=10,
        )
        ws.column_dimensions[letra].width = min(ancho + 4, 48)
    ws.freeze_panes = "A2"


def _respuesta_xlsx(wb: Workbook, nombre_archivo: str) -> StreamingResponse:
    buffer = BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    return StreamingResponse(
        buffer,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f'attachment; filename="{nombre_archivo}"'},
    )


def _nombre_curso(db: Session, id_curso: int) -> str:
    curso = curso_repository.get(db, id_curso)
    return curso.nombre if curso else f"Curso #{id_curso}"


def _nombre_nivel_de_curso(db: Session, id_curso: int) -> str:
    curso = curso_repository.get(db, id_curso)
    if not curso:
        return "-"
    nivel = nivel_repository.get(db, curso.id_nivel)
    return nivel.nombre if nivel else "-"


# ==================== ALUMNOS Y DOCENTES ====================
@router.get(
    "/alumnos.xlsx", summary="Nómina de alumnos y docentes en Excel (solo ADMIN)"
)
def exportar_alumnos_docentes(
    token: str | None = Query(None),
    db: Session = Depends(get_db),
):
    _autorizar_admin(db, token)
    wb = Workbook()
    wb.remove(wb.active)

    alumnos = alumno_repository.get_multi(db, limit=2000)
    filas_alumnos = [
        [
            a.legajo,
            a.apellido,
            a.nombre,
            a.dni,
            _nombre_curso(db, a.id_curso),
            _nombre_nivel_de_curso(db, a.id_curso),
            a.estado_alumno,
            a.telefono or "-",
            a.correo_electronico or "-",
        ]
        for a in sorted(alumnos, key=lambda x: (x.apellido, x.nombre))
    ]
    _agregar_hoja(
        wb,
        "Alumnos",
        [
            "Legajo",
            "Apellido",
            "Nombre",
            "DNI",
            "Curso",
            "Nivel",
            "Estado",
            "Teléfono",
            "Correo",
        ],
        filas_alumnos,
    )

    profesores = profesor_repository.get_multi(db, limit=2000)
    filas_profesores = [
        [
            p.legajo,
            p.apellido,
            p.nombre,
            p.dni,
            p.especialidad,
            p.estado,
            p.telefono,
            p.correo_electronico,
        ]
        for p in sorted(profesores, key=lambda x: (x.apellido, x.nombre))
    ]
    _agregar_hoja(
        wb,
        "Docentes",
        [
            "Legajo",
            "Apellido",
            "Nombre",
            "DNI",
            "Especialidad",
            "Estado",
            "Teléfono",
            "Correo",
        ],
        filas_profesores,
    )
    return _respuesta_xlsx(wb, "nomina_alumnos_docentes.xlsx")


# ==================== CURSOS Y MATERIAS ====================
@router.get(
    "/cursos-materias.xlsx",
    summary="Cursos, materias y docentes a cargo en Excel (solo ADMIN)",
)
def exportar_cursos_materias(
    token: str | None = Query(None),
    db: Session = Depends(get_db),
):
    _autorizar_admin(db, token)
    wb = Workbook()
    wb.remove(wb.active)

    cursos = curso_repository.get_multi(db, limit=1000)
    filas_cursos = []
    for c in cursos:
        nivel = nivel_repository.get(db, c.id_nivel)
        cantidad = len(alumno_repository.get_by_curso(db, c.id_curso))
        filas_cursos.append(
            [c.id_curso, c.nombre, nivel.nombre if nivel else "-", cantidad]
        )
    _agregar_hoja(wb, "Cursos", ["ID", "Curso", "Nivel", "Alumnos"], filas_cursos)

    asignaciones = academico_service.get_asignaciones(db)
    filas_asig = [
        [a.materia, a.curso, a.nivel, a.profesor, a.legajo_profesor]
        for a in sorted(asignaciones, key=lambda x: (x.nivel, x.curso, x.materia))
    ]
    _agregar_hoja(
        wb,
        "Asignaciones",
        ["Materia", "Curso", "Nivel", "Docente", "Legajo docente"],
        filas_asig,
    )

    materias = materia_repository.get_multi(db, limit=1000)
    _agregar_hoja(
        wb,
        "Materias",
        ["ID", "Materia"],
        [[m.id_materia, m.nombre] for m in sorted(materias, key=lambda x: x.nombre)],
    )
    return _respuesta_xlsx(wb, "cursos_y_materias.xlsx")


# ==================== DEPORTES ====================
@router.get(
    "/deportes.xlsx", summary="Reporte integral de deportes en Excel (solo ADMIN)"
)
def exportar_deportes(
    token: str | None = Query(None),
    db: Session = Depends(get_db),
):
    _autorizar_admin(db, token)
    wb = Workbook()
    wb.remove(wb.active)

    grupos = servicios_escolares_service.get_grupos_deporte(db)
    inscripciones = servicios_escolares_service.get_inscripciones_deporte(db)
    inscriptos_por_grupo = {}
    for i in inscripciones:
        inscriptos_por_grupo[i.id_grupo] = inscriptos_por_grupo.get(i.id_grupo, 0) + 1

    _agregar_hoja(
        wb,
        "Grupos",
        ["Grupo", "Disciplina", "Nivel", "Horario", "Inscriptos"],
        [
            [
                g.id_grupo,
                g.deporte,
                g.nivel,
                g.horario,
                inscriptos_por_grupo.get(g.id_grupo, 0),
            ]
            for g in grupos
        ],
    )
    _agregar_hoja(
        wb,
        "Inscripciones",
        ["Alumno", "Disciplina", "Fecha de inscripción"],
        [
            [i.alumno, i.deporte, i.fecha_inscripcion.strftime("%d/%m/%Y")]
            for i in inscripciones
        ],
    )
    return _respuesta_xlsx(wb, "reporte_deportes.xlsx")


# ==================== TRANSPORTE ====================
@router.get(
    "/transporte.xlsx", summary="Hoja de ruta de transporte en Excel (solo ADMIN)"
)
def exportar_transporte(
    token: str | None = Query(None),
    db: Session = Depends(get_db),
):
    _autorizar_admin(db, token)
    wb = Workbook()
    wb.remove(wb.active)

    recorridos = servicios_escolares_service.get_recorridos(db)
    inscripciones = servicios_escolares_service.get_inscripciones_transporte(db)
    activos_por_recorrido = {}
    for i in inscripciones:
        if i.estado == "ACTIVO":
            activos_por_recorrido[i.id_recorrido] = (
                activos_por_recorrido.get(i.id_recorrido, 0) + 1
            )

    _agregar_hoja(
        wb,
        "Recorridos",
        ["ID", "Recorrido", "Pasajeros activos"],
        [
            [
                r.id_recorrido,
                r.descripcion,
                activos_por_recorrido.get(r.id_recorrido, 0),
            ]
            for r in recorridos
        ],
    )
    _agregar_hoja(
        wb,
        "Pasajeros",
        ["Alumno", "Recorrido", "Fecha de alta", "Estado"],
        [
            [i.alumno, i.recorrido, i.fecha_alta.strftime("%d/%m/%Y"), i.estado]
            for i in sorted(inscripciones, key=lambda x: x.alumno)
        ],
    )
    return _respuesta_xlsx(wb, "transporte_pasajeros.xlsx")


# ==================== SOLICITUDES DE ADMISIÓN ====================
@router.get(
    "/solicitudes.xlsx",
    summary="Listado de solicitudes de admisión en Excel (solo ADMIN)",
)
def exportar_solicitudes(
    token: str | None = Query(None),
    db: Session = Depends(get_db),
):
    _autorizar_admin(db, token)
    wb = Workbook()
    wb.remove(wb.active)

    solicitudes = comunidad_service.get_solicitudes(db)
    _agregar_hoja(
        wb,
        "Solicitudes",
        [
            "N°",
            "Fecha",
            "Postulante",
            "DNI",
            "Nivel",
            "Tutor",
            "Teléfono",
            "Correo",
            "Estado",
            "Observaciones",
        ],
        [
            [
                f"#SOL-{s.fecha.year}-{s.id_solicitud:04d}",
                s.fecha.strftime("%d/%m/%Y"),
                s.postulante,
                s.dni,
                s.nivel,
                s.tutor,
                s.telefono or "-",
                s.correo or "-",
                s.estado,
                s.observaciones or "",
            ]
            for s in solicitudes
        ],
    )
    return _respuesta_xlsx(wb, "solicitudes_admision.xlsx")


# ==================== RESUMEN GENERAL ====================
@router.get("/resumen.xlsx", summary="Resumen general de gestión en Excel (solo ADMIN)")
def exportar_resumen(
    token: str | None = Query(None),
    db: Session = Depends(get_db),
):
    _autorizar_admin(db, token)
    wb = Workbook()
    wb.remove(wb.active)

    comunidad = comunidad_service.resumen(db)
    transporte = [
        i
        for i in inscripcion_transporte_repository.get_multi(db, limit=1000)
        if i.estado == "ACTIVO"
    ]
    comedor = [
        i
        for i in inscripcion_comedor_repository.get_multi(db, limit=1000)
        if i.estado == "ACTIVO"
    ]

    metricas = [
        ["Alumnos activos", len(alumno_repository.get_activos(db, limit=1000))],
        ["Docentes activos", len(profesor_repository.get_activos(db, limit=1000))],
        ["Cursos", len(curso_repository.get_multi(db, limit=500))],
        ["Niveles educativos", len(nivel_repository.get_multi(db, limit=100))],
        ["Grupos de deporte", len(grupo_deporte_repository.get_multi(db, limit=300))],
        ["Pasajeros de transporte activos", len(transporte)],
        ["Inscriptos al comedor activos", len(comedor)],
        ["Solicitudes de admisión pendientes", comunidad["solicitudes_pendientes"]],
        ["Mensajes de contacto sin leer", comunidad["mensajes_no_leidos"]],
        ["Opiniones pendientes de moderación", comunidad["opiniones_pendientes"]],
    ]
    _agregar_hoja(wb, "Resumen", ["Métrica", "Valor"], metricas)
    return _respuesta_xlsx(wb, "resumen_general.xlsx")
